"""将 API文档.md 转换为 Word 文档"""
import re
import os
import sys
sys.path.insert(0, r'D:\Lib\site-packages')
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_shading(cell, color_hex):
    """设置单元格背景色"""
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color_hex,
    })
    shading.append(shd)

def add_code_block(doc, code_text):
    """添加代码块"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    # 设置中文字体
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')

def add_table_from_rows(doc, headers, rows):
    """从表头和数据行创建表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # 表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ''
        p = cell.paragraphs[0]
        run = p.add_run(header.strip())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_shading(cell, '4080D0')

    # 数据行
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= len(headers):
                break
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(cell_text.strip())
            run.font.size = Pt(10)
            # 隔行变色
            if row_idx % 2 == 0:
                set_cell_shading(cell, 'F0F4F8')

    doc.add_paragraph()  # 表格后空行

def parse_markdown_table(lines, start_idx):
    """解析 markdown 表格，返回 (headers, rows, next_index)"""
    # 表头行
    header_line = lines[start_idx].strip()
    headers = [h.strip() for h in header_line.split('|')[1:-1]]

    # 跳过分隔行
    next_idx = start_idx + 2
    rows = []
    while next_idx < len(lines) and lines[next_idx].strip().startswith('|'):
        row_data = [c.strip() for c in lines[next_idx].strip().split('|')[1:-1]]
        rows.append(row_data)
        next_idx += 1

    return headers, rows, next_idx

def convert_md_to_docx(md_path, docx_path):
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 设置标题样式
    for i in range(1, 5):
        heading_style = doc.styles[f'Heading {i}']
        heading_style.font.name = '微软雅黑'
        heading_style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
        if i == 1:
            heading_style.font.size = Pt(22)
            heading_style.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)
        elif i == 2:
            heading_style.font.size = Pt(16)
            heading_style.font.color.rgb = RGBColor(0x20, 0x4B, 0x8B)
        elif i == 3:
            heading_style.font.size = Pt(14)
            heading_style.font.color.rgb = RGBColor(0x20, 0xB2, 0xAA)
        else:
            heading_style.font.size = Pt(12)
            heading_style.font.color.rgb = RGBColor(0x40, 0x80, 0xD0)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    code_buffer = []

    while i < len(lines):
        line = lines[i].rstrip('\n')

        # 代码块处理
        if line.strip().startswith('```'):
            if in_code_block:
                # 结束代码块
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
                in_code_block = False
            else:
                # 开始代码块
                in_code_block = True
                code_buffer = []
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # 标题
        if line.startswith('# ') and not line.startswith('## '):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        # 水平线
        if line.strip() == '---':
            # 添加分页符
            doc.add_page_break()
            i += 1
            continue

        # 表格
        if line.strip().startswith('|') and i + 1 < len(lines) and '---' in lines[i + 1]:
            headers, rows, next_i = parse_markdown_table(lines, i)
            add_table_from_rows(doc, headers, rows)
            i = next_i
            continue

        # 引用块
        if line.startswith('> '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.color.rgb = RGBColor(0x64, 0x74, 0x8b)
            run.italic = True
            i += 1
            continue

        # 列表项
        if re.match(r'^[\-\*] ', line):
            text = re.sub(r'^[\-\*] ', '', line).strip()
            # 处理加粗和行内代码
            p = doc.add_paragraph(style='List Bullet')
            add_formatted_text(p, text)
            i += 1
            continue

        # 有序列表
        if re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line).strip()
            p = doc.add_paragraph(style='List Number')
            add_formatted_text(p, text)
            i += 1
            continue

        # 空行
        if not line.strip():
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_formatted_text(p, line.strip())
        i += 1

    # 处理末尾未关闭的代码块
    if code_buffer:
        add_code_block(doc, '\n'.join(code_buffer))

    doc.save(docx_path)
    print(f"Word文档已生成: {docx_path}")

def add_formatted_text(paragraph, text):
    """处理加粗、行内代码等格式"""
    # 先处理行内代码 `code`
    parts = re.split(r'(`[^`]+`)', text)
    for part in parts:
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        elif '**' in part:
            # 处理加粗
            bold_parts = re.split(r'(\*\*[^*]+\*\*)', part)
            for bp in bold_parts:
                if bp.startswith('**') and bp.endswith('**'):
                    run = paragraph.add_run(bp[2:-2])
                    run.bold = True
                else:
                    run = paragraph.add_run(bp)
        else:
            paragraph.add_run(part)

if __name__ == '__main__':
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_path = os.path.join(project_root, 'API文档.md')
    docx_path = os.path.join(project_root, 'API文档.docx')
    convert_md_to_docx(md_path, docx_path)
